from docx import Document


def extract_text_from_docx(file):
    """
    Extract text from a DOCX file.

    Text is extracted from:
    1. Normal paragraphs
    2. Table cells

    Empty paragraphs and empty table cells are ignored.
    """

    document = Document(file)

    extracted_parts = []

    # -------------------------
    # EXTRACT PARAGRAPHS
    # -------------------------

    for paragraph in document.paragraphs:

        text = paragraph.text.strip()

        if text:
            extracted_parts.append(
                text
            )

    # -------------------------
    # EXTRACT TABLES
    # -------------------------

    for table in document.tables:

        for row in table.rows:

            row_text = []

            for cell in row.cells:

                text = cell.text.strip()

                if text:
                    row_text.append(
                        text
                    )

            if row_text:

                extracted_parts.append(
                    " | ".join(row_text)
                )

    # -------------------------
    # COMBINE TEXT
    # -------------------------

    final_text = "\n".join(
        extracted_parts
    )

    return {
        "text": final_text,
        "has_text": bool(
            final_text.strip()
        )
    }